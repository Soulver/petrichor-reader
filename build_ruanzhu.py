# -*- coding: utf-8 -*-
"""Generate software copyright identification materials (docx)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "软著申请材料"
SOFTWARE = "基于动态字体与Unicode私有使用区码位映射的文本防采集器"
VERSION = "V1.0"
HEADER = f"{SOFTWARE} {VERSION}"

SOURCE_FILES = [
    "packages/font-tool/src/policy.js",
    "packages/font-tool/src/mapText.js",
    "packages/font-tool/src/buildFont.js",
    "packages/font-tool/src/cli.js",
    "packages/font-tool/python/remap_subset.py",
    "packages/font-tool/scripts/download-master-font.mjs",
    "packages/reader-api/src/config.js",
    "packages/reader-api/src/revision.js",
    "packages/reader-api/src/richText.js",
    "packages/reader-api/src/chapters.js",
    "packages/reader-api/src/rateLimit.js",
    "packages/reader-api/src/tickets.js",
    "packages/reader-api/src/remoteArticle.js",
    "packages/reader-api/src/obfuscate.js",
    "packages/reader-api/src/server.js",
    "packages/reader-api/src/index.js",
    "packages/reader-web/src/server.js",
    "packages/reader-web/src/index.js",
    "packages/reader-web/public/read.html",
    "packages/reader-web/public/read.css",
    "packages/reader-web/public/read.js",
    "packages/reader-web/public/embed.js",
    "examples/serve.mjs",
    "examples/host.html",
    "packages/font-tool/test/mapText.test.js",
    "packages/font-tool/test/pipeline.test.js",
    "packages/reader-api/test/richText.test.js",
    "packages/reader-api/test/revision.test.js",
    "packages/reader-api/test/obfuscate.test.js",
    "packages/reader-api/test/remoteArticle.test.js",
    "packages/reader-api/test/api.test.js",
    "packages/reader-web/test/page.test.js",
]


def set_run_font(run, east="宋体", ascii_font="Times New Roman", size=12, bold=False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = ascii_font
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east)


def set_paragraph_format(p, before=0, after=6, line=22, first_line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = Pt(line)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)


def add_page_number(paragraph):
    run1 = paragraph.add_run("第 ")
    set_run_font(run1, size=9)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r2 = paragraph.add_run()
    r2._element.append(fld1)
    r3 = paragraph.add_run()
    r3._element.append(instr)
    r4 = paragraph.add_run()
    r4._element.append(fld2)
    run5 = paragraph.add_run(" 页")
    set_run_font(run5, size=9)


def setup_header_footer(section, extra=""):
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = hp.add_run(HEADER + extra)
    set_run_font(run, east="宋体", size=9)
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp)


def setup_page(doc, top=2.54, bottom=2.54, left=2.54, right=2.54):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    setup_header_footer(section)
    return section


def add_title(doc, text, size=16):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, before=6, after=10, line=28)
    run = p.add_run(text)
    set_run_font(run, east="黑体", ascii_font="Arial", size=size, bold=True)


def add_heading_cn(doc, text, level=1):
    size = 14 if level == 1 else 13 if level == 2 else 12
    p = doc.add_paragraph()
    set_paragraph_format(p, before=10, after=6, line=24)
    run = p.add_run(text)
    set_run_font(run, east="黑体", ascii_font="Arial", size=size, bold=True)


def add_body(doc, text, indent=True):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=6, line=22, first_line=0.74 if indent else 0)
    run = p.add_run(text)
    set_run_font(run, size=12)


def add_bullet(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=3, line=22, first_line=0)
    p.paragraph_format.left_indent = Cm(0.74)
    run = p.add_run("• " + text)
    set_run_font(run, size=12)


def add_pre(doc, text):
    for line in text.splitlines() or [""]:
        p = doc.add_paragraph()
        set_paragraph_format(p, before=0, after=0, line=16, first_line=0)
        run = p.add_run(line if line else " ")
        set_run_font(run, east="宋体", ascii_font="Consolas", size=9)


def spec_paragraphs():
    return [
        ("title", "计算机软件著作权登记说明书"),
        ("h1", "引言"),
        ("p", f"本文档旨在为“{SOFTWARE}”（以下简称“本软件”）申请中华人民共和国国家版权局计算机软件著作权登记，提供系统性的技术说明与功能证明。作为鉴别材料之文档部分，本文档将从问题域界定、体系结构、核心算法、安全边界与工程实现等层面，对本软件的独创性表达予以完整呈现。"),
        ("p", "本软件系作者针对公开万维网超文本环境中，长文本正文易被自动化程序经由标记语言源码、文档对象模型文本接口及应用程序接口响应体直接提取这一技术事实，独立构思并实现的原创性软件作品。其名称中“Unicode私有使用区码位映射”所指，为将可读字符之Unicode标量值映射至Unicode私有使用区（Private Use Area，PUA）及必要时之补充私有使用区，并与动态生成的OpenType子集字体之字符到字形映射表（cmap）形成共生交付关系。文档重点阐述如下方面："),
        ("bullet", "问题域与设计目标：论证何以前端脚本加密、交互抑制及简单样式替换不足以切断“码位同一性”，从而必须将明文留置服务端，对外仅输出私有区码位与配套动态字体。"),
        ("bullet", "功能构成：说明字体流水线、短时阅读凭证、富文本消毒与文本节点编码、阅读器表现层、双层水印、宿主嵌入协议以及按更新时间失效的修订缓存如何形成闭环。"),
        ("bullet", "算法与协议：给出私有区编号函数、cmap重写与按需子集化、源隔离之跨源资源共享策略、超文本结构保持条件下的可见文本收集，以及修订键构造方法。"),
        ("bullet", "独创性归纳：阐明本软件并非对通用阅读器或字体库的简单拼装，而在于“编码与字体必须同时交付、结构与文本必须分离处理、权限沿浏览器源递减”这一方法组合。"),
        ("h1", "一、软件概述与开发背景"),
        ("h2", "1.1 软件全称、简称与版本号"),
        ("p", f"软件全称：{SOFTWARE}"),
        ("p", "软件简称：文本防采集器"),
        ("p", f"版本号：{VERSION}"),
        ("p", "软件分类建议：应用软件／网络应用软件。开发的硬件环境为通用微型计算机；运行的硬件环境为通用微型计算机或云服务器；开发所用操作系统为GNU/Linux或Microsoft Windows；主要开发语言为JavaScript（ECMAScript模块）、Python 3、HTML5与CSS3；运行平台为Node.js 18及以上运行时及现代浏览器；权利取得方式为原始取得，开发方式为独立开发。"),
        ("h2", "1.2 开发背景与目的"),
        ("p", "背景分析。在当代网络文学及个人站点长文发布场景中，章节正文普遍以超文本标记语言或JavaScript对象简谱明文写入浏览器。采集程序只需复现浏览器已完成的解析过程，即可经由innerText、textContent或接口响应获得与读者视觉感知一致的自然语言文本。既有应对手段存在结构性缺陷："),
        ("bullet", "交互层抑制（禁止选择、禁止上下文菜单）仅作用于用户代理的交互通道，对无头浏览器及源码读取无效。"),
        ("bullet", "传输层或脚本内对称加密之密钥与算法必然随脚本下发，敌手复现同一计算过程即可还原明文，其保密性不成立。"),
        ("bullet", "仅改写层叠样式表而不改写字体cmap时，文档中的Unicode标量值仍与自然语言文本同一，复制操作即可得到原文。"),
        ("bullet", "将防护逻辑耦合于内容管理前端仓库，使算法生命周期与站点业务绑定，难以独立演进、独立部署与独立主张著作权。"),
        ("p", "中文正文之字体授权构成额外约束。对禁止修改字形或禁止再分发之字库擅自改写cmap，将构成许可违约。故本软件之母版字体选用SIL Open Font License 1.1许可之Noto Sans SC静态字形，并在生成子集时强制改写名称表，禁止继续使用Noto、Source等保留家族名。"),
        ("p", "开发目的。本软件旨在实现一套与宿主站点解耦的文本防采集子系统，使“人眼可读、机器直接采集不可得”成为默认对外形态，具体包括："),
        ("bullet", "明文仅短暂存在于阅读器服务端进程及受控稿件源；对外接口不返回汉字区等自然语言码位。"),
        ("bullet", "按章节可见字符构建单射，将非空白字符映射至以U+E000起始之BMP私有区，越界后延拓至U+F0000起之补充私有区。"),
        ("bullet", "依映射结果裁剪母版并重写cmap，输出WOFF2；空白类字符按ECMAScript正则\\s原样保留，以维持排版骨架。"),
        ("bullet", "超文本结构与文本节点分离处理：标签及安全属性得以保留，仅对文本节点及可能泄漏原文之alt等属性实施编码。"),
        ("bullet", "阅读器页面须成功加载专用字体后方可呈现正文；失败时仅提示“内容无法显示”，禁止回退明文。"),
        ("bullet", "宿主仅引入嵌入脚本与占位节点；跨窗口消息只传递数值高度，不传递正文。正文接口之CORS白名单不含宿主源。"),
        ("bullet", "远程稿件以更新时间（及必要之内容摘要）构造修订号，使字体缓存与正文缓存随稿件变更失效。"),
        ("h2", "1.3 主要功能简介"),
        ("p", f"本软件{VERSION}围绕上述目的，实现下列相互耦合之功能："),
        ("bullet", "动态子集字体生成：据本章可见字符建立“原字符→私有区码位”映射，经Python侧fontTools完成子集化、cmap替换、名称表改写与WOFF2封装。"),
        ("bullet", "短时阅读凭证：以密码学安全伪随机数签发票证，绑定章节标识、会话标识与字体资源标识，默认生存期三百秒。"),
        ("bullet", "富文本消毒与节点编码：解析并白名单化常见排版标签，剔除脚本及危险协议，仅编码可见文本。"),
        ("bullet", "阅读器表现层：字号调节、深浅主题、失败关闭；偏好存于本地存储且不持久化章节文本。"),
        ("bullet", "阅读区水印：浅层（近背景色）与深层（可见但不夺目）两套倾斜重复纹理，模式为关闭、仅浅、仅深、双层。"),
        ("bullet", "宿主嵌入：扫描data-petrichor-reader节点，以iframe打开阅读页，并以MutationObserver兼容单页应用挂载时序。"),
        ("bullet", "远程稿件适配与修订失效：按文章编号拉取超文本，剥离危险标记后进入编码流水线；每次解析回源，字体缓存键含更新时间。"),
        ("bullet", "访问控制：源白名单、可选站点密钥、按客户端地址之滑动窗口限流。"),
        ("h1", "二、系统运行环境要求"),
        ("h2", "2.1 硬件环境"),
        ("p", "本软件以通用计算设备为载体，无需图形处理器加速。建议配置如下：中央处理器为x86_64或ARM64多核；内存不低于四吉字节，以容纳Node.js进程、Python子进程及章节字体缓冲；存储预留不少于二吉字节，母版字体单独占用数十兆字节；生产环境须保证阅读器源可被访客用户代理访问。"),
        ("h2", "2.2 软件环境"),
        ("bullet", "操作系统：GNU/Linux、Microsoft Windows 10及以上，或macOS。"),
        ("bullet", "运行时：Node.js 18及以上；Python 3.10及以上，依赖fonttools及brotli以写出WOFF2。"),
        ("bullet", "用户代理：Chromium系浏览器（推荐）、Gecko、WebKit；不以旧式Internet Explorer为目标平台。"),
        ("bullet", "反向代理（可选）：Nginx，将/v1/与静态资源分别反代至本机回环地址上的接口进程与前端进程。"),
        ("h2", "2.3 开发语言与工具"),
        ("p", "服务端与工具链采用JavaScript与Python 3；浏览器端为原生HTML5、CSS3与JavaScript，不依赖前端框架。字体处理采用fontTools.ttLib与fontTools.subset。版本管理采用Git。宿主示例为独立静态页；生产宿主可为既有内容站点，仅增加占位节点与脚本地址配置。"),
        ("h1", "三、系统功能模块详解"),
        ("h2", "3.1 系统功能结构"),
        ("p", "本软件在仓库内划分为三个可独立验证之软件包，并与宿主保持权限边界：工具链包font-tool唯一允许可逆映射与字体改写；接口包reader-api持有明文之短暂驻留权；表现层包reader-web不得包含逆映射表。宿主站点仅引用embed.js。"),
        ("pre", "宿主站点\n  embed.js + [data-petrichor-reader][data-chapter-id]\n        │ iframe（阅读器前端源）\n        ▼\nreader-web  /read  /embed.js\n        │ CORS 仅允许阅读器源\n        ▼\nreader-api  换票 / 取编码正文 / 取 WOFF2\n        │ 修订键 = 标识 + 更新时间或内容摘要\n        ▼\nfont-tool   映射策略 + remap_subset.py\n        │\n        ▼\n稿件源：本地样章或远程文章接口"),
        ("p", "图1 功能与部署结构示意图。建议申请人于独立阅读页、宿主嵌入页分别截取运行界面，作为操作说明附图。"),
        ("h2", "3.2 字符映射与动态子集字体模块"),
        ("p", "该模块构成本软件防护能力之根基。其并不试图在浏览器中隐藏字符串，而是改变字符串所使用的Unicode标量值，同时提供一张仅在这些新标量值上绘制正确轮廓的字体。于是：在加载该字体的排版引擎中，读者获得正常字形；在未加载该字体或仅提取码位的采集程序中，得到的是私有区字符，不能直接当作自然语言正文使用。"),
        ("p", "策略判定。对输入逐一枚举用户可见字符。若字符匹配ECMAScript之Unicode空白（含普通空格、制表符、CR/LF、不间断空格、表意空格等），则作为直通字符保留。其余可见字符——汉字、拉丁字母、数字、标点及符号——均分配私有区码位。"),
        ("p", "稳定编号。同一章节内，首次出现的需映射字符按出现顺序获得自零起始之索引。设BMP私有区长度为六千四百，索引k若小于该长度，则码位为U+E000+k，否则为U+F0000+(k−6400)。该策略保证常规模用字落于BMP，同时为极端长文预留扩展。"),
        ("p", "编码与失败断言。编码函数对直通字符原样写出，其余替换为映射后的私有区字符。服务端在外发前检查编码结果之可见文本是否仍含CJK统一汉字区码位；若仍含汉字，则视为流水线失败，避免半编码泄漏。"),
        ("p", "子集与cmap重写。将“原码位→目标码位”列表连同母版路径、新家族名、PostScript名及直通码位以JSON经标准输入交付Python脚本。脚本以TTFont读取母版，收集所需原码位与豆腐字形（U+25A1），以Subsetter裁剪，再以format 4或format 12子表替换cmap：直通码位仍指向空白等原字形，映射项将目标私有区码位指向原汉字字形；母版缺字则指向豆腐或.notdef。名称表写入前缀为pr-sess-之家族名及PrSess前缀之PostScript名，flavor设为woff2后保存。映射表不得打包进阅读器前端。"),
        ("h2", "3.3 富文本消毒与结构保持模块"),
        ("p", "既有内容管理系统常以超文本存储章节，并含加粗、斜体、列表、链接与图片。若将整段标记字符串纳入映射，则标签名本身亦被编入私有区，结构必然崩解。本模块因此将“结构”与“文本”分离。"),
        ("p", "规范化。若输入不具标签形态，则按空行分段、按单行换行转写为段落与换行元素。若输入为超文本，则进入消毒解析。"),
        ("p", "白名单。允许之元素包括p、br、strong、b、em、i、u、blockquote、h1至h3、ul、ol、li、a、img。span、div等包裹元素予以解包以保留子孙。script、style、iframe、object、embed、link、meta、noscript予以丢弃并跳过其内部。"),
        ("p", "属性约束。图像源仅允许http、https或站点根相对路径，拒绝javascript:及data:。锚点地址仅允许http、https、mailto:或根相对路径，并强制rel为noopener noreferrer。alt、title作为可见文本之延伸，参与收集与编码，以免属性侧信道泄漏原文。Quill对齐类名被改写为本软件可识别之pr-align-*，其余类名与事件属性一律剔除。"),
        ("p", "可见文本收集。遍历文本节点及alt、title，其连接结果用于构建映射表，从而保证映射域与将被编码的字符集合一致。仅含图像而无文字之章节仍视为有效载荷。"),
        ("h2", "3.4 阅读凭证、会话与字体句柄模块"),
        ("p", "若编码正文与字体统一资源定位符可被任意第三方长期缓存或猜测，则采集者仍可建立稳定流水线。故引入短时阅读凭证。签发时校验可选站点密钥，确认章节可解析后，以crypto.randomBytes生成票证与字体资源标识，写入内存表并记录过期时刻。消费章节接口必须持票；通过后执行混淆，并将WOFF2缓冲与同一过期时刻写入字体表。每次查询修剪过期项。字体标识与票证字符串相分离，以降低日志侧信道风险。"),
        ("h2", "3.5 阅读器接口模块"),
        ("p", "接口以Node.js原生http实现，主要路径包括：元数据查询、换票、持票取编码正文、按标识取WOFF2。若请求携带Origin且不在允许列表中，立即拒绝。允许列表默认仅含阅读器前端源。按客户端地址统计每分钟请求数，逾限返回过载。章节解析优先读取本地样章夹具；未命中且配置了远程文章接口时，以表单编码提交标识并解析返回对象。文首、正文、文尾超文本经消毒后合并。失败时返回字体不可用等错误码，前端不得展示疑似明文。"),
        ("h2", "3.6 修订识别与缓存失效模块"),
        ("p", "既有实现曾以章节标识为唯一缓存键，导致作者修订后读者仍获得旧正文与旧字体，直至进程重启。本软件以修订号切断该粘滞。"),
        ("p", "远程稿件在每次解析时回源，不再以标识为键长期复用旧正文；并发解析以进行中的承诺对象合并，避免重复打源。修订号优先取文章更新时间，其次创建时间；二者皆空时，以标题与正文之SHA-1摘要前十六位十六进制作为内容修订，以免后台未填写时间戳时缓存永不失效。"),
        ("p", "字体缓存键为“章节标识＋修订号”。命中则直接返回既有字形串与WOFF2缓冲；未命中则重新映射并生成字体，同时驱逐该标识下之其他修订，防止内存中堆积历史版本。磁盘侧仍以章节标识为文件名覆盖写出。浏览器侧字体统一资源定位符含票内随机字体标识，不与修订号耦合，亦不会把旧文件长期热链至新稿。"),
        ("h2", "3.7 阅读器页面与水印模块"),
        ("p", "阅读页为唯一被允许持票拉取正文的用户界面。其换票、取章、校验fontUrl、fontFamily与bodyGlyphs之合取条件，并以FontFace（display为block）加载WOFF2成功后，将标题写入文本节点，将正文以消毒后的超文本写入容器，字体家族作用于该容器及其后代。任一步骤失败均转入统一失败文案。"),
        ("p", "水印配置于本版本写死于客户端脚本，包含模式与文案。模式取值：off、light、dark、both。浅水印颜色贴近当前主题背景，深层水印提高不透明度但仍保持可读性。二者以可缩放矢量图形数据URI生成倾斜文字瓦片并重复铺满阅读主区域，深层相对浅层错开背景定位，以免完全重叠。指针事件关闭，以免遮挡选择与滚动。主题切换时依当前层叠变量重算填充色。"),
        ("p", "嵌入适配。若页面处于iframe中，为根元素增加embedded类以取消最小视口高度。通过postMessage向父窗口发送类型为petrichor-reader:resize之消息，载荷仅含高度。ResizeObserver在排版变化时再次上报。父窗口脚本校验event.origin必须等于阅读器源，且仅当event.source对应当前iframe时改写高度。"),
        ("h2", "3.8 宿主嵌入模块"),
        ("p", "embed.js以立即执行函数封装，自currentScript.src推导阅读器源，从而在开发端口、反代域名与生产域名之间保持可移植。其查找带有data-petrichor-reader之元素，读取data-chapter-id，创建或复用iframe。脚本早于节点到达时，于DOMContentLoaded之后启动，并以MutationObserver观察子树与属性，覆盖单页应用路由切换。对外暴露扫描函数供宿主在数据到达后主动触发。宿主不得再将文章超文本以不安全方式写入自身文档。目录标题、作者、评论等元数据可继续明文存在。"),
        ("h1", "四、技术架构与核心算法实现"),
        ("h2", "4.1 总体技术架构"),
        ("p", "本软件采用“工具链／接口层／表现层／宿主层”四段结构，权限沿远离明文的方向递减。接口层使用内存映射保存票证与字体，不引入外部键值存储，以降低本版本运维面；进程重启后未过期票证失效，客户端可重新换票。生成字体之计算成本高于稿件回源，故回源每次执行，而昂贵之子集化结果按修订号复用。"),
        ("h2", "4.2 核心算法一：面向排版的私有区映射"),
        ("p", "采集程序所依赖者，为码位同一性：若文档中的字符与自然语言文本使用同一Unicode标量值，则文本接口之输出即原文。本算法切断该同一性，但保留排版所需的“码位→字形”关系。"),
        ("p", "形式化描述。设章节可见文本为字符序列c1…cn。谓词W(c)表示c为空白。σ为从非空白字符到自然数的单射，σ(c)等于该字符首次出现之次序（自零计）。定义π(k)：若k<6400，则π(k)=0xE000+k，否则π(k)=0xF0000+(k−6400)。编码E(c)=c若W(c)，否则E(c)=chr(π(σ(c)))。服务端断言E作用于可见文本后不匹配汉字区正则，方可外发。"),
        ("p", "与固定偏移或凯撒替换之区别在于：编号由本章用字集合之出现顺序决定，而非全局常数；且必须与cmap改写同时发布，否则用户代理只能显示缺失字形。采集者若欲还原，需下载WOFF2并解析cmap与轮廓表，已超出“直接读取文档文本”之威胁模型。本软件不声称防御光学识别、截图与专业字体逆向。"),
        ("h2", "4.3 核心算法二：cmap重写与按需子集化"),
        ("p", "若仅替换码位而不改字体，读者将看到方框；若下发完整母版，则体积不可接受且扩大对照面。故必须只保留本章用到的原字形，并将其挂到新的私有区码位上。"),
        ("p", "子集阶段关闭 hinting 与版式特性，desubroutinize 以避免CFF子例程在改写cmap后失效，并丢弃DSIG。映射阶段对每一对(src,dst)，若src存在于子集后的cmap，则新表以dst指向该字形；若母版本无该字，则指向豆腐。直通码位自子集cmap复制。随后按最大码位选择BMP之format 4或补充平面之format 12，分别写入Microsoft与Unicode平台，以免不同用户代理选取不同子表而不一致。Node进程不在其内解析OpenType，而以子进程调用Python，任务经标准输入传递。"),
        ("h2", "4.4 核心算法三：结构保持条件下的超文本编码"),
        ("p", "设消毒后之文档为树T。函数Vis(T)为全部文本节点及指定属性值之连接。映射表M=BuildMap(Vis(T)∪标题)。编码遍历T，对文本节点t以M逐字替换，对alt、title作同样替换，再序列化为标记语言。标签名、允许之属性名及安全地址字面量不进入M。于是读者所见仍为段落、强调与图像，而文档对象模型中的字符数据为私有区码位。"),
        ("p", "该算法在防护与可用性之间取得可证明的边界：结构信息（如何分段、何为强调、图像位于何处）被保留；语义载体（字是什么）被替换。图像之源地址本非字符正文，予以保留；此举不扩大“正文被直接采集”之威胁面，亦不以图像防盗为目标。"),
        ("h2", "4.5 核心协议：短时票、锁源与修订键"),
        ("p", "防护不仅是编码问题，亦是“谁被允许触发编码”的问题。正文跨源策略白名单不含宿主源，宿主只负责嵌入iframe。时序为：嵌入脚本设置iframe地址→阅读页换票→持票取章→加载字体→写入节点→回传高度。票在生存期内有效，过期后须重换；字体统一资源定位符同步过期。"),
        ("p", "修订键R=updatedAt若非空，否则R=H(title∥body)之前缀。字体缓存键K=id∥R。当且仅当K命中，复用生成结果。该协议保证：标识不变而更新时间变化时，必然未命中；更新时间缺失而正文变化时，摘要变化亦未命中。"),
        ("h1", "五、操作使用说明"),
        ("h2", "5.1 启动服务"),
        ("p", "于安装Node.js与Python依赖并放置母版字体后，分别启动接口进程与前端进程。默认接口监听本机三千九百八十端口，阅读页为该机三千九百八十一端口之/read。应以用户代理打开，勿以文件协议打开宿主示例。"),
        ("p", "图2 独立阅读页运行效果（请申请人自行截取：正文可读、水印铺满阅读区、工具条可见）。"),
        ("h2", "5.2 阅读、主题与水印"),
        ("p", "打开阅读页后，系统自动换票并加载本章字体。字号增减限制在十六至三十二像素。主题在浅色纸色与深色背景之间切换，水印颜色随主题重算。水印模式与文案于本版本在客户端常量中配置，取值见第三章。"),
        ("p", "图3 深色主题及双水印效果（请截取）。"),
        ("h2", "5.3 验证防采集是否生效"),
        ("p", "在阅读页打开开发者工具元素面板，标题与正文之文本节点应为私有区字符而非原文。网络面板中章节响应不应出现可读汉字正文。禁用该自定义字体后，页面应变为方框或乱码，且不得自动改回明文。若章节含加粗或图像，元素树中应仍能观察到相应元素，但其文本数据仍为私有区码位。"),
        ("p", "图4 开发者工具中正文为私有区码位之截图（请截取，可对样章原文打码）。"),
        ("h2", "5.4 嵌入宿主与稿件修订"),
        ("p", "最小接入为在宿主引入阅读器源上的embed.js，并放置带有data-petrichor-reader与data-chapter-id之占位节点。作者于内容系统中修订章节并保存后，读者重新打开或刷新阅读页，阅读器将回源取得新的更新时间并在修订变化时重新生成字体，无需重启服务进程。"),
        ("p", "图5 宿主章节页嵌入后之整体版面（请截取）。"),
        ("h2", "5.5 生产部署要点"),
        ("p", "阅读器公共基址必须为访客可访问之统一资源定位符。若站点启用传输层安全，阅读器亦须使用同源方案，否则用户代理将拦截混合内容。接口进程可仅监听回环地址，由反向代理对外。母版字体因体积不纳入版本库，部署时须单独放置，缺失将导致全文无法显示。"),
        ("h1", "六、软件独创性、特点与总结"),
        ("h2", "6.1 技术创新点"),
        ("p", "面向文档对象模型采集威胁模型的“编码与字体共生”方案。本软件不把机密置于浏览器脚本，而使浏览器成为只认字形、不认码位的排版终端。私有区映射与cmap重写必须同时交付，缺一不可，构成方法上的核心独创性。"),
        ("p", "结构保持的超文本编码。区别于将章节坍缩为纯文本的简化实现，本软件在白名单消毒之后仅替换文本节点，使防护与排版表达并存。该分离并非对通用消毒库的套用，而是服务于既定映射不变量：映射域等于可见文本集合。"),
        ("p", "以更新时间为轴的生成物失效。将昂贵的子集化结果与廉价的回源解耦：每次读取均可观察到最新修订号，仅当修订变化时才重新承担字体生成代价。更新时间缺省时回退至内容摘要，避免“无时间戳即永不过期”之退化。"),
        ("p", "锁源、短时票与iframe隔离的系统化组合。即便编码正确，若宿主页能直接取接口或把明文写进自身文档，防护仍告失败。本软件以跨源白名单、票绑定字体句柄、嵌入层只传高度、失败关闭四条规则同时收紧接口权限与文档权限。"),
        ("p", "与内容站点解耦的嵌入协议。防护实现独立成仓，宿主仅增加标记属性，使本软件可作为独立作品主张著作权，而不与特定博客业务代码混同。"),
        ("h2", "6.2 软件特点"),
        ("bullet", "威胁模型清晰：阻止经由标记语言、文档文本接口与章节接口直接获得可读正文；不夸大为防御截图或光学识别。"),
        ("bullet", "失败安全：字体或接口异常时不回退明文。"),
        ("bullet", "排版可用：支持段落、强调、列表、链接与图像等常见长文结构。"),
        ("bullet", "修订一致：稿件更新后生成字体随修订号失效。"),
        ("bullet", "实现可审计：核心路径为有限之JavaScript与一份Python脚本，无闭源字体引擎。"),
        ("h2", "6.3 总结"),
        ("p", f"综上所述，“{SOFTWARE}”是作者独立设计、自主开发的原创性软件成果。它针对公开网页长文易被自动化采集之问题，将Unicode私有区码位映射、OpenType动态子集化、超文本结构保持、短时访问凭证、源隔离与修订失效组合成可部署的文本防采集系统，并提供标准嵌入方式接入既有站点。"),
        ("p", "本软件在映射与字体共生、结构与文本分离、修订键缓存以及宿主—阅读器解耦等方面具有可说明的独创性，功能闭环完整，适于作为计算机软件著作权登记之技术说明依据。"),
        ("p", "补充说明甲：接口错误码包括源不允许、无效JSON、站点密钥不匹配、章节不存在、票无效、字体不可用、限流、字体未找到等。JSON响应禁止存储缓存。"),
        ("p", "补充说明乙：前端静态服务将路径限制于公共目录，拒绝目录穿越；仅允许GET与HEAD；超文本发出前替换接口基址占位。"),
        ("p", "补充说明丙：水印与站点密钥于本版本为客户端常量，可在后续版本改为配置下发而不改变核心映射方法。"),
        ("p", "补充说明丁：明确非目标包括付费墙、登录门禁、多租户后台、禁止复制、前端分组密码及将阅读器注入宿主文档对象模型。"),
        ("p", "补充说明戊：术语表。私有区指Unicode私有使用区；cmap指OpenType字符到字形索引表；WOFF2指网络开放字体格式第二版；修订号指更新时间或内容摘要；夹具指开发用样章。"),
        ("p", f"补充说明己：版本记录。{VERSION}实现阅读闭环、宿主嵌入、富文本结构保持、双层水印及按更新时间之缓存失效。后续功能变更应另行办理变更或补充登记，并保持页眉中的软件名称与版本号与申请表一致。"),
    ]


def write_spec(path: Path):
    doc = Document()
    setup_page(doc)
    for kind, text in spec_paragraphs():
        if kind == "title":
            add_title(doc, text)
            sub = doc.add_paragraph()
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_format(sub, after=12, line=22)
            r = sub.add_run(f"{SOFTWARE}\n版本号：{VERSION}")
            set_run_font(r, east="黑体", size=12, bold=True)
        elif kind == "h1":
            add_heading_cn(doc, text, 1)
        elif kind == "h2":
            add_heading_cn(doc, text, 2)
        elif kind == "bullet":
            add_bullet(doc, text)
        elif kind == "pre":
            add_pre(doc, text)
        else:
            add_body(doc, text)
    doc.save(path)


def collect_source_lines():
    lines = []
    for rel in SOURCE_FILES:
        fp = ROOT / rel
        if not fp.exists():
            continue
        text = fp.read_text(encoding="utf-8", errors="replace").replace("\r\n", "\n")
        lines.append(f"// ===== FILE: {rel} =====")
        for raw in text.split("\n"):
            s = raw.replace("\t", "    ")
            while True:
                lines.append(s[:100])
                s = s[100:]
                if not s:
                    break
        lines.append("")
    return lines


def write_source(path: Path):
    all_lines = collect_source_lines()
    per = 50
    total = len(all_lines)
    pages_all = (total + per - 1) // per
    if pages_all <= 60:
        selected = all_lines
        note = f"程序鉴别材料（全文，共 {total} 行，约 {pages_all} 页）"
    else:
        selected = all_lines[: per * 30] + ["", "// ===== 以下为源程序连续后 30 页 =====", ""] + all_lines[-per * 30 :]
        note = f"程序鉴别材料（一般交存：前 30 页 + 后 30 页；全文约 {pages_all} 页 / {total} 行）"

    doc = Document()
    section = setup_page(doc, top=2.0, bottom=2.0, left=2.0, right=2.0)
    for p in section.header.paragraphs:
        p.clear()
    run = section.header.paragraphs[0].add_run(HEADER + "  源程序")
    set_run_font(run, size=9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(title, after=8, line=20)
    r = title.add_run(note)
    set_run_font(r, east="黑体", size=12, bold=True)

    warn = doc.add_paragraph()
    set_paragraph_format(warn, after=8, line=18)
    wr = warn.add_run("排版说明：每页不少于 50 行；页眉软件名称与版本号与申请表保持一致；有无字母 V 以申请表为准。打印或另存 PDF 时请保持纵向 A4、不要缩放。")
    set_run_font(wr, size=10)

    for i, line in enumerate(selected, 1):
        p = doc.add_paragraph()
        set_paragraph_format(p, before=0, after=0, line=13.2, first_line=0)
        run = p.add_run(f"{i:04d}| " + (line if line else " "))
        set_run_font(run, east="宋体", ascii_font="Consolas", size=8)
        if i % per == 0 and i < len(selected):
            runb = p.add_run()
            brk = OxmlElement("w:br")
            brk.set(qn("w:type"), "page")
            runb._element.append(brk)

    doc.save(path)
    return total, pages_all


def write_checklist(path: Path, src_lines: int, src_pages: int):
    path.write_text(
        f"""软件著作权登记材料清单（对照中国版权保护中心「所需文件」）

软件全称：{SOFTWARE}
软件简称：文本防采集器
版本号：{VERSION}

【官方要求摘要】
申请文件应包括：软件著作权登记申请表、软件的鉴别材料、相关证明文件。纵向排版，文字从左向右。
页眉软件版本号须与申请表一致（有无字母 V 以申请表为准）。

一、本文件夹已生成（请转为 PDF 后上传）
1. 软件说明书.docx —— 文档鉴别材料。不足 60 页则提交全文。
2. 源程序鉴别材料.docx —— 每页 50 行。当前约 {src_lines} 行、约 {src_pages} 页；不足 60 页则提交全文。

二、须在登记系统在线完成
1. R11 申请表；在线打印申请确认签章页，不得改内容/格式/打印比例，签章后上传 PDF。
2. 身份证明（个人身份证正反面，或单位营业执照副本）。

三、申请表填写建议（请按实际情况修改）
软件全称：{SOFTWARE}
软件简称：文本防采集器
版本号：V1.0
软件分类：应用软件 → 网络应用软件（以系统下拉为准）
开发完成日期：【填写实际完成日】
发表状态：未发表（若已公开网站则选已发表并填首次发表日）
开发方式：独立开发
权利取得方式：原始取得
编程语言：JavaScript、Python、HTML、CSS
源程序量：约 {src_lines} 行（以鉴别材料为准）
开发目的：在网页环境中展示长文时，通过动态子集字体与Unicode私有区码位映射，降低正文被经由标记语言、文档对象模型文本接口及章节接口直接采集的风险，并与宿主站点解耦部署。
主要功能和技术特点：见说明书第一、三、四、六章。
运行支撑环境 / 开发环境：见说明书第二章。

四、请自行截图插入说明书第五章后转 PDF
1. 独立阅读页正文可读，水印可见或按配置呈现。
2. 深色主题。
3. 开发者工具 Elements：正文为私有区码位。
4. Network 中章节响应无可读汉字正文。
5. 宿主章节页整体（标题/阅读区/评论），宿主 DOM 无正文超文本。
6. （可选）含加粗或插图的章节，结构仍在但文字为私有区码位。
7. （可选）禁用自定义字体后出现方框或“内容无法显示”。
""",
        encoding="utf-8",
    )


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    spec_path = OUT / "软件说明书.docx"
    src_path = OUT / "源程序鉴别材料.docx"
    list_path = OUT / "材料清单与填表参考.txt"
    write_spec(spec_path)
    n, pages = write_source(src_path)
    write_checklist(list_path, n, pages)
    print("wrote", spec_path)
    print("wrote", src_path, "lines", n, "pages", pages)
    print("wrote", list_path)


if __name__ == "__main__":
    main()
